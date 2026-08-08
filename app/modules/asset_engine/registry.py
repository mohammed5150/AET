"""Asset registry import from XLSX exports (SDS-004).

Reads an airport asset-registry workbook into normalized ``Asset`` records
with type classification (§5.1) and circuit derivation from asset names
(§5.2). File-level failures raise :class:`InputError`; row-level failures
are skipped and counted (§6).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from app.core.errors import InputError
from app.core.logging import StructuredLogger, get_logger
from app.models.asset import Asset
from app.models.geometry import parse_coordinate
from app.models.project import SourceInput

_CIRCUIT_PREFIX = re.compile(r"^([A-Za-z]+)")

DEFAULT_TYPE_RULES: tuple[tuple[str, str], ...] = (
    (r"^AGL PIT$", "agl-pit"),
    (r"INSET|ELEV", "light-fitting"),
    # Fitting shorthands without an INSET/ELEV token (SDS-004 §5.1 rev. A):
    # taxiway edge, stop bar, taxiway centreline, lead-in, runway guard.
    (r"^(TWE|STB)-in", "light-fitting"),
    (r"^TCL-", "light-fitting"),
    (r"^LIL_", "light-fitting"),
    (r"^(RGL|STL)$", "light-fitting"),
    (r"^SGN|^SG$", "sign"),
    (r"^RRM", "rrm"),
    (r"NBASE|EBASE|^Base|^CRBLK", "base"),
    (r"^Lightpoles$", "lightpole"),
    (r"^CCR-", "ccr"),
    (r"^High Mast", "high-mast"),
    (r"Streetlight", "streetlight"),
    (r"^Apron Stand$", "apron-floodlight"),
    (r"Obstruction", "obstruction-light"),
    (r"^TRL", "traffic-light"),
    (r"^WDI$", "wdi"),
    # Auxiliary crossing/access equipment tracked outside the core AGL
    # taxonomy ("extra assets"): crossing barriers/controls (TRA), flasher
    # control units (FCU), runway access equipment (RRS), LVO equipment.
    (r"^(TRA|FCU|RRS|LVO)", "extra-asset"),
)

# Name-based fallback when the class column is unusable ("Generic", "*U",
# damaged exports): the circuit family embedded in the asset name still
# identifies what the asset is (SDS-004 §5.3).
FAMILY_TYPE_FALLBACK: tuple[tuple[str, str], ...] = (
    ("TCC", "light-fitting"),
    ("TEC", "light-fitting"),
    ("SBC", "light-fitting"),
    ("LIC", "light-fitting"),
    ("RCC", "light-fitting"),
    ("REC", "light-fitting"),
    ("RGC", "light-fitting"),
    ("SGC", "sign"),
    ("HH", "agl-pit"),
)

_PASSTHROUGH_COLUMNS: tuple[tuple[str, str], ...] = (
    ("id", "registry_id"),
    ("identifier", "identifier"),
    ("serialNumber", "serial_number"),
    ("specialInstructions", "special_instructions"),
    ("genericText1", "generic_text1"),
    ("genericText2", "generic_text2"),
    ("genericText3", "generic_text3"),
)


@dataclass(frozen=True, slots=True)
class RegistryColumnMap:
    """Header names in the registry export (defaults match the AUH export)."""

    name: str = "name"
    asset_class: str = "assetClass"
    main_area: str = "mainArea"
    sub_area: str = "subArea"
    utm_zone: str = "utmZone"
    utm_e: str = "utmE"
    utm_n: str = "utmN"

    def required(self) -> tuple[str, ...]:
        return (self.name, self.asset_class, self.main_area)


@dataclass(frozen=True, slots=True)
class RegistryImport:
    """Result of reading one registry workbook."""

    assets: list[Asset] = field(default_factory=list)
    skipped_rows: int = 0


def classify_asset_type(
    asset_class: str,
    rules: tuple[tuple[str, str], ...] = DEFAULT_TYPE_RULES,
    name: str = "",
) -> str:
    """Map an assetClass value to an asset type (SDS-004 §5.1, §5.3).

    When class-based rules yield no match and a ``name`` is given, the
    circuit family parsed from the name is tried as a fallback.
    """
    for pattern, asset_type in rules:
        if re.search(pattern, asset_class, flags=re.IGNORECASE):
            return asset_type
    if name:
        family = derive_circuit(name).get("circuit_family", "")
        for prefix, asset_type in FAMILY_TYPE_FALLBACK:
            if family.startswith(prefix):
                return asset_type
    return "other"


def derive_circuit(name: str) -> dict[str, str]:
    """Parse circuit attributes from an asset name prefix (SDS-004 §5.2)."""
    prefix = re.split(r"[.\-]", name, maxsplit=1)[0].strip()
    match = _CIRCUIT_PREFIX.match(prefix)
    if not match:
        return {}
    family = match.group(1).upper()
    circuit = prefix.upper()
    derived = {"circuit_family": family}
    if circuit != family:
        derived["circuit"] = circuit
    return derived


class XlsxAssetRegistryReader:
    """Reads XLSX registry exports into normalized assets (SDS-004 §3-§5)."""

    sheet_name = "Assets"

    def __init__(
        self,
        columns: RegistryColumnMap | None = None,
        type_rules: tuple[tuple[str, str], ...] = DEFAULT_TYPE_RULES,
        logger: StructuredLogger | None = None,
    ) -> None:
        self._columns = columns or RegistryColumnMap()
        self._type_rules = type_rules
        self._logger = logger or get_logger("asset-registry")

    def read(self, source: SourceInput) -> RegistryImport:
        """Read the registry file registered as ``source``."""
        path = Path(source.path)
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except FileNotFoundError as exc:
            raise InputError(
                f"Registry file not found: {path}",
                remediation="Check the file path and re-import the registry.",
            ) from exc
        except Exception as exc:  # noqa: BLE001 - openpyxl raises broadly
            raise InputError(
                f"Failed to read registry workbook {path}: {exc}",
                remediation="Verify the file is a valid .xlsx export.",
            ) from exc
        try:
            sheet = (
                workbook[self.sheet_name]
                if self.sheet_name in workbook.sheetnames
                else workbook.worksheets[0]
            )
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                raise InputError(
                    f"Registry sheet '{sheet.title}' in {path} is empty",
                    remediation="Export the registry with a header row.",
                )
            index = {
                str(cell).strip(): position
                for position, cell in enumerate(header)
                if cell is not None
            }
            missing = [
                column for column in self._columns.required() if column not in index
            ]
            if missing:
                raise InputError(
                    f"Registry {path} is missing required column(s): "
                    f"{', '.join(missing)}",
                    remediation=(
                        "Adjust the export or provide a RegistryColumnMap "
                        "matching its headers."
                    ),
                )
            return self._read_rows(rows, index, source)
        finally:
            workbook.close()

    def _read_rows(self, rows, index: dict[str, int], source) -> RegistryImport:
        columns = self._columns

        def cell(row: tuple, column: str) -> str:
            position = index.get(column)
            if position is None or position >= len(row):
                return ""
            value = row[position]
            return "" if value is None else str(value).strip()

        assets: list[Asset] = []
        skipped = 0
        for row_number, row in enumerate(rows, start=2):
            name = cell(row, columns.name)
            if not name:
                skipped += 1
                self._logger.warn(
                    f"Skipped registry row {row_number}: empty name",
                    stage="asset-derivation",
                    project_id=source.project_id,
                )
                continue
            asset_class = cell(row, columns.asset_class)
            attributes = {
                "asset_class": asset_class,
                "main_area": cell(row, columns.main_area),
                "registry_input_id": source.input_id,
            }
            if sub_area := cell(row, columns.sub_area):
                attributes["sub_area"] = sub_area
            for column, key in _PASSTHROUGH_COLUMNS:
                if value := cell(row, column):
                    attributes[key] = value
            attributes.update(derive_circuit(name))
            raw_zone = cell(row, columns.utm_zone)
            raw_easting = cell(row, columns.utm_e)
            raw_northing = cell(row, columns.utm_n)
            location = parse_coordinate(raw_easting, raw_northing, zone=raw_zone)
            if location is None:
                # Keep what the export actually said: an ordinate we cannot
                # use must stay visible for correction, not vanish silently
                # (SDS-009 §5.3).
                if raw_zone:
                    attributes["utm_zone"] = raw_zone
                if raw_easting or raw_northing:
                    attributes["utm_unparsed"] = f"{raw_easting}|{raw_northing}"
                    self._logger.warn(
                        f"Row {row_number} ({name}): unusable UTM ordinates "
                        f"{raw_easting!r}/{raw_northing!r}",
                        stage="asset-derivation",
                        project_id=source.project_id,
                    )
            assets.append(
                Asset(
                    project_id=source.project_id,
                    snapshot_id="",
                    asset_type=classify_asset_type(
                        asset_class, self._type_rules, name=name
                    ),
                    name=name,
                    location=location,
                    attributes=attributes,
                )
            )
        return RegistryImport(assets=assets, skipped_rows=skipped)
